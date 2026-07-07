# -*- coding: utf-8 -*-
"""
階段1解析：待評估資訊/20260707-清潔隊部資訊 → outputs/20260707-清潔隊部資訊/01_中繼檔.xlsx

本腳本針對「清潔隊備勤室非都市用地對應國土功能分區彙整表」的表格版面
（序號/縣市/所屬單位/清潔隊名/鄉鎮市區/地址/地號/都市用地/土地分區/使用地目/國土功能分區）。
其他案子的輸入版面不同時，由 Claude 對話中另寫對應的解析驅動（設計文件 §5.1）。

輸出中繼檔一列一個地號，欄位依設計文件 §5.1；「解析信心=低」的列以黃底標示，
供使用者人工核對（兩段式流程的人工把關點）。

執行：python -X utf8 scripts/parse_stage1.py
設計文件：docs/superpowers/specs/2026-07-07-大量圖資轉換工具-design.md §5.1
"""

from __future__ import annotations

import re
import sys
import unicodedata
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import pandas as pd

from parse_helpers import (
    LandNoParseError,
    landno_display,
    normalize_county,
    parse_landno_verbose,
)

CASE_NAME = "20260707-清潔隊部資訊"
PDF_PATH = REPO_ROOT / "待評估資訊" / CASE_NAME / f"{CASE_NAME}v1.pdf"
OUT_DIR = REPO_ROOT / "outputs" / CASE_NAME
OUT_XLSX = OUT_DIR / "01_中繼檔.xlsx"

# 一個範圍展開的上限，超過視為解析可疑（避免把地號誤判成範圍炸出大量列）
RANGE_EXPAND_LIMIT = 60

# ---------------------------------------------------------------------------
# 逐案特例（人工判讀後登錄；均會寫入備註欄讓使用者核對）
#   pre:          地號欄字串前處理
#   note:         整案備註（附加到該案所有地號列）
#   low:          整案標低信心
#   parcel_notes: {地號原文: (備註, 是否低信心)}
# ---------------------------------------------------------------------------
OVERRIDES: dict[int, dict] = {
    9: {
        "note": "「(地先)」為地籍慣用註記（該地號臨接之未登錄地），非正式地號",
    },
    10: {
        "pre": lambda s: s.replace(".", "、"),
        "note": "原文以「.」分隔（13-1.80.80-2），已視為頓號分隔",
    },
    27: {
        "parcel_notes": {
            "77": ("國土功能分區欄提及「地號79」，與地號欄之77不一致，請確認是77或79", True),
        },
    },
    48: {
        "pre": lambda s: s.replace("彰化市安溪", "彰化市安溪段"),
        "note": "原文「彰化市安溪518…」缺「段」字，依彰化市地籍判定為安溪段",
    },
    55: {
        "note": "「2-1~7-1」之展開方式（2-1、3-1…7-1）為推測，請確認",
    },
    57: {
        "pre": lambda s: s.replace("186-193", "186~193"),
        "note": "「186-193」依國土功能分區欄逐號列舉（186~193共8筆）判定為連續範圍；"
        "使用分區欄原文跨列且疑有誤字（特地農業區特定農業區）",
    },
    62: {
        "pre": lambda s: s.replace("番路村番路段", "番路段"),
        "note": "原文段名前綴「番路村」為村名（番路鄉轄下），非鄉鎮名，已移除",
        "parcel_notes": {
            "63-116": ("「63-116」可能為63-116地號、也可能為63~116範圍，請確認", True),
        },
    },
    82: {
        "parcel_notes": {
            "1-3": ("「1-3」可能為1-3地號、也可能為1~3範圍，請確認", True),
        },
    },
    86: {
        "note": "段藉索引查無高樹鄉「東振段」，僅有「東振新」段；原表國土功能分區欄亦註明系統查無，請確認段名",
        "low": True,
    },
    90: {
        "pre": lambda s: s.replace("-資收", ""),
        "note": "原文「328地號-資收」含用途註記（資收），已移除",
    },
}

# ---------------------------------------------------------------------------
# 地號欄解析
# ---------------------------------------------------------------------------

# 段/小段 標頭：[鄉鎮市區]? + X段 + [Y小段]?
# sec 首字不允許數字，避免把前面殘留的地號數字吃進段名
_SEG_RE = re.compile(
    r"(?P<town>[一-鿿]{1,3}[鄉鎮市區])?"
    r"(?P<sec>[一-鿿][一-鿿\d]{0,9}?段)"
    r"(?P<sub>[一-鿿]{1,8}?小段)?"
)
_LIST_SEP = re.compile(r"[、,，;；及]+")
_PAREN_RE = re.compile(r"[（(]([^）)]*)[）)]")
_COUNT_HINT_RE = re.compile(r"等(\d+)筆(?:地號)?")


def _clean_cell(s: str) -> str:
    s = unicodedata.normalize("NFKC", s or "")
    return re.sub(r"[\s　]+", "", s)


def _expand_range(a: str, b: str):
    """展開「A~B」範圍。回傳 (清單[(母,子)], 是否有歧義)；無法展開時 raise LandNoParseError。"""
    pa = parse_landno_verbose(a)
    pb = parse_landno_verbose(b)
    if isinstance(pa, LandNoParseError):
        raise pa
    if isinstance(pb, LandNoParseError):
        raise pb
    (ma, ca), (mb, cb) = pa, pb

    if ma == mb and ca != cb:
        # 628-10~628-17：同母號的子號範圍
        lo, hi = sorted((int(ca), int(cb)))
        pairs = [(ma, f"{i:04d}") for i in range(lo, hi + 1)]
        ambiguous = False
    elif ca == "0000" and cb == "0000":
        # 1849~1855：母號範圍
        lo, hi = sorted((int(ma), int(mb)))
        pairs = [(f"{i:04d}", "0000") for i in range(lo, hi + 1)]
        ambiguous = False
    else:
        # 2-1~7-1：母號遞增、子號固定 —— 展開方式為推測，標低信心
        lo, hi = sorted((int(ma), int(mb)))
        pairs = [(f"{i:04d}", ca) for i in range(lo, hi + 1)]
        ambiguous = True

    if len(pairs) > RANGE_EXPAND_LIMIT:
        raise LandNoParseError(f"範圍展開超過{RANGE_EXPAND_LIMIT}筆，疑非範圍", f"{a}~{b}")
    return pairs, ambiguous


def parse_parcel_cell(cell: str, seq: int) -> list[dict]:
    """解析地號欄字串，回傳地號 dict 清單。

    每筆：{town, sec, sub, raw, mother, child, annotation, conf_low, notes, expected_count}
    town/sec/sub 為原文（sec 含「段」、sub 含「小段」字尾）；解析失敗的片段
    以 mother=None 輸出並附原因，讓使用者在中繼檔上看得到、可手動補正。
    """
    ov = OVERRIDES.get(seq, {})
    s = _clean_cell(cell)
    if "pre" in ov:
        s = ov["pre"](s)

    expected_count = None
    m = _COUNT_HINT_RE.search(s)
    if m:
        expected_count = int(m.group(1))
        s = _COUNT_HINT_RE.sub("", s)

    # 「地號」一律視為分隔符（真實段名不含「地號」二字），避免相鄰段名誤切
    # （如「尚德段2-1~7-1地號明德段529-2」）；「地號和大漢段」的「和」一併清掉
    s = s.replace("地號", "、").replace("、和", "、")

    seg_matches = [m for m in _SEG_RE.finditer(s)]
    results: list[dict] = []

    for i, seg in enumerate(seg_matches):
        start = seg.end()
        end = seg_matches[i + 1].start() if i + 1 < len(seg_matches) else len(s)
        parcels_text = s[start:end]

        for token in _LIST_SEP.split(parcels_text):
            token = token.strip()
            if not token:
                continue
            token = re.sub(r"號$", "", token)
            if not token:
                continue

            annotations = _PAREN_RE.findall(token)
            token_clean = _PAREN_RE.sub("", token).strip()
            if not token_clean:
                continue

            base = {
                "town": seg.group("town"),
                "sec": seg.group("sec"),
                "sub": seg.group("sub"),
                "expected_count": expected_count,
            }

            annotation = ""
            conf_low = False
            notes: list[str] = []
            for ann in annotations:
                if "部分" in ann:
                    annotation = ann  # 部分 / 部分使用
                elif "地先" in ann:
                    notes.append("原文註記「(地先)」")
                    conf_low = True
                else:
                    notes.append(f"原文含未識別註記「({ann})」")
                    conf_low = True

            if "~" in token_clean:
                a, b = token_clean.split("~", 1)
                try:
                    pairs, ambiguous = _expand_range(a, b)
                except LandNoParseError as exc:
                    results.append({**base, "raw": token, "mother": None, "child": None,
                                    "annotation": annotation, "conf_low": True,
                                    "notes": notes + [f"範圍無法展開：{exc.reason}"]})
                    continue
                for mother, child in pairs:
                    r_notes = list(notes) + [f"由範圍「{token_clean}」展開"]
                    r_low = conf_low or ambiguous
                    results.append({**base, "raw": token, "mother": mother, "child": child,
                                    "annotation": annotation, "conf_low": r_low,
                                    "notes": r_notes})
            else:
                parsed = parse_landno_verbose(token_clean)
                if isinstance(parsed, LandNoParseError):
                    results.append({**base, "raw": token, "mother": None, "child": None,
                                    "annotation": annotation, "conf_low": True,
                                    "notes": notes + [f"地號無法解析：{parsed.reason}"]})
                else:
                    mother, child = parsed
                    results.append({**base, "raw": token, "mother": mother, "child": child,
                                    "annotation": annotation, "conf_low": conf_low,
                                    "notes": list(notes)})

    return results


# ---------------------------------------------------------------------------
# PDF 表格列抽取
# ---------------------------------------------------------------------------


def extract_site_rows(pdf_path: Path) -> list[dict]:
    """抽出每個案場（序號）一筆的原始欄位；跨頁/跨列續行自動併回前一筆。"""
    sites: list[dict] = []
    with pdfplumber.open(pdf_path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            for table in page.extract_tables():
                for row in table:
                    cells = [(c or "").replace("\n", "") for c in row]
                    # 表頭列（序號…）與子表頭列（土地分區/使用地目）跳過
                    if cells[0] == "序號" or (not cells[0] and "土地分區" in cells):
                        continue
                    if cells[0].strip().isdigit():
                        sites.append({
                            "seq": int(cells[0]),
                            "page": page_no,
                            "county": cells[1],
                            "agency": cells[2],
                            "site_name": cells[3],
                            "town": cells[4],
                            "address": cells[5],
                            "parcel_cell": cells[6],
                            "urban": cells[7],
                            "zone_in": cells[8],
                            "landuse_in": cells[9],
                            "fz_in": cells[10],
                        })
                    elif any(cells) and sites:
                        # 跨列續行：非空欄位補到前一筆對應欄位
                        keymap = {1: "county", 2: "agency", 3: "site_name", 4: "town",
                                  5: "address", 6: "parcel_cell", 7: "urban",
                                  8: "zone_in", 9: "landuse_in", 10: "fz_in"}
                        for idx, key in keymap.items():
                            if idx < len(cells) and cells[idx]:
                                sites[-1][key] = (sites[-1][key] + cells[idx]).strip()
    return sites


# ---------------------------------------------------------------------------
# 鄉鎮名清理與段藉驗證
# ---------------------------------------------------------------------------


def _tai_norm(s):
    return s.replace("台", "臺") if isinstance(s, str) else s


def clean_town(town: str, county: str) -> str:
    """鄉鎮欄清理：異體字「巿」→「市」、去除縣市名前綴（如 桃園市觀音區 → 觀音區）。"""
    t = _clean_cell(town).replace("巿", "市")
    county_n = normalize_county(county) or ""
    for prefix in (county_n, county_n.replace("臺", "台")):
        if prefix and t.startswith(prefix):
            t = t[len(prefix):]
    return t


def load_current_section_index() -> pd.DataFrame:
    """讀取由地籍 shp 掃出的現行段/小段索引（scripts/build_section_index.py 產出）。

    注意：段藉.shp 為舊制段名索引、不含重測後新段名（如深坑區賴仲段），
    不能當驗證依據；現行索引即比對階段實際使用的地籍 shp 屬性，權威一致。
    """
    csv_path = REPO_ROOT / "config" / "section_index.csv"
    df = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig")
    for col in ("縣市", "鄉鎮", "段", "小段"):
        df[f"{col}_t"] = df[col].map(_tai_norm)
    return df


def _town_base(t: str) -> str:
    return t[:-1] if t and t[-1] in "鄉鎮市區" else t


def _is_subsequence(a: str, b: str) -> bool:
    it = iter(b)
    return all(ch in it for ch in a)


def town_matches(index_town: str, input_town: str) -> bool:
    """鄉鎮名寬鬆比對。

    容忍兩類已知差異（比對階段以段碼為準，不受影響）：
    - 行政區改制字尾不同：員林鎮(地籍舊制) vs 員林市
    - 地籍 Big5 轉檔缺字：蘆區(缺「竹」) vs 蘆竹區
    """
    if not isinstance(index_town, str) or not isinstance(input_town, str):
        return False
    a, b = _tai_norm(index_town), _tai_norm(input_town)
    if a == b:
        return True
    ab, bb = _town_base(a), _town_base(b)
    if ab == bb:
        return True
    # 缺字容忍：索引名為輸入名的子序列且僅少一字
    return len(ab) >= 1 and len(bb) - len(ab) == 1 and _is_subsequence(ab, bb)


def _suggest_sections(sec_df: pd.DataFrame, county_n: str, town_n: str, sec_name: str) -> str:
    """查無段名時，找同縣市內「僅差一字」的相近段名（優先同鄉鎮）供人工確認。"""

    def similar(a: str, b: str) -> bool:
        if abs(len(a) - len(b)) == 1:
            longer, shorter = (a, b) if len(a) > len(b) else (b, a)
            return longer.startswith(shorter) or longer.endswith(shorter)
        if len(a) == len(b):
            return sum(1 for x, y in zip(a, b) if x != y) == 1
        return False

    cand = sec_df.loc[sec_df["縣市_t"] == county_n, ["鄉鎮", "段_t"]].drop_duplicates()
    hits = cand[cand["段_t"].map(lambda n: isinstance(n, str) and similar(n, sec_name))]
    if hits.empty:
        return ""
    same_town = hits[hits["鄉鎮"].map(lambda t: town_matches(t, town_n))]
    picked = same_town if not same_town.empty else hits
    return "、".join(
        f"{r.鄉鎮}「{r.段_t}」段" for r in picked.head(3).itertuples()
    )


def validate_section(sec_df: pd.DataFrame, county: str, town: str, sec: str, sub):
    """對現行段/小段索引驗證 (縣市, 鄉鎮, 段, 小段) 是否存在。

    以「縣市＋段名」為主鍵、鄉鎮名寬鬆比對（地籍鄉鎮欄有舊制名與缺字）。
    回傳 (ok, 備註, 是否低信心)。ok 且備註非空 = 資訊性提示（不降信心）。
    """
    county_n = normalize_county(county)
    town_n = _tai_norm(town)
    sec_name = _tai_norm(re.sub(r"段$", "", sec))
    sub_name = _tai_norm(re.sub(r"小段$", "", sub)) if sub else None

    cand = sec_df.loc[
        (sec_df["縣市_t"] == county_n) & (sec_df["段_t"] == sec_name)
    ]
    if cand.empty:
        hint = _suggest_sections(sec_df, county_n, town_n, sec_name)
        msg = f"地籍索引查無 {county_n}「{sec_name}」段"
        if hint:
            msg += f"；相近段名：{hint}，請確認是否為原表誤字"
        return False, msg, True

    towns = [t for t in cand["鄉鎮"].dropna().unique()]
    matched_towns = [t for t in towns if town_matches(t, town_n)]

    note = ""
    if len(matched_towns) == 1:
        picked = matched_towns[0]
        if _tai_norm(picked) != town_n:
            note = f"地籍鄉鎮欄記載為「{picked}」（舊制/缺字），比對以段碼為準不受影響"
        subset = cand[cand["鄉鎮"] == picked]
    elif len(matched_towns) == 0:
        other = "、".join(towns[:4])
        return False, f"「{sec_name}」段在地籍中位於{other}，與表列鄉鎮（{town_n}）不符，請確認", True
    else:
        other = "、".join(matched_towns[:4])
        return False, f"「{sec_name}」段在{other}皆有同名段，無法判定，請確認", True

    if sub_name is not None:
        hit = subset[subset["小段_t"] == sub_name]
        if hit.empty:
            subs = "、".join(str(x) for x in subset["小段"].dropna().unique()[:8])
            return False, f"地籍索引查無「{sec_name}」段之「{sub_name}」小段（現有小段：{subs or '無'}）", True
        return True, note, False

    # 未指定小段：若該段所有記錄都有小段，表示輸入漏了小段
    if len(subset) > 0 and subset["小段"].notna().all():
        subs = "、".join(str(x) for x in subset["小段"].dropna().unique()[:6])
        return False, f"「{sec_name}」段之下均為小段（{subs}…），原文未指定小段", True
    return True, note, False


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

COLUMNS = [
    "案場ID", "案場名稱", "縣市", "鄉鎮區", "段", "小段", "地號原文",
    "宗地母", "宗地子", "部分註記", "輸入使用分區", "輸入使用地目",
    "輸入國土功能分區", "解析信心", "備註", "所屬單位", "地址",
]


def main() -> None:
    sites = extract_site_rows(PDF_PATH)
    print(f"抽出案場數：{len(sites)}")

    sec_df = load_current_section_index()

    out_rows: list[list] = []
    issues: list[str] = []

    for site in sites:
        seq = site["seq"]
        ov = OVERRIDES.get(seq, {})
        county = _clean_cell(site["county"])
        town_default = clean_town(site["town"], county)

        parcels = parse_parcel_cell(site["parcel_cell"], seq)
        if not parcels:
            issues.append(f"[{seq:03d}] 地號欄無法解析出任何地號：{site['parcel_cell']}")

        # 「等N筆」核對
        expected = parcels[0]["expected_count"] if parcels else None
        count_ok = expected is None or len(parcels) == expected
        if not count_ok:
            issues.append(f"[{seq:03d}] 展開{len(parcels)}筆 ≠ 原文「等{expected}筆」")

        site_notes: list[str] = []
        if ov.get("note"):
            site_notes.append(ov["note"])
        urban = _clean_cell(site["urban"])
        if urban and urban != "無":
            site_notes.append(f"都市用地欄：{urban}")
        if not count_ok:
            site_notes.append(f"地號展開{len(parcels)}筆與原文「等{expected}筆」不符，請確認")

        # 段/小段驗證（同案場相同組合只查一次）
        sec_check_cache: dict[tuple, tuple[bool, str]] = {}

        for p in parcels:
            town = clean_town(p["town"], county) if p["town"] else town_default
            notes = list(site_notes)
            conf_low = p["conf_low"] or ov.get("low", False) or (not count_ok)
            notes.extend(p["notes"])

            if p["town"] and clean_town(p["town"], county) != town_default:
                notes.append(f"地號欄鄉鎮（{town}）與表列鄉鎮（{town_default}）不同，以地號欄為準")

            pn = ov.get("parcel_notes", {}).get(p["raw"])
            if pn:
                notes.append(pn[0])
                conf_low = conf_low or pn[1]

            sec, sub = p["sec"], p["sub"]
            key = (town, sec, sub)
            if key not in sec_check_cache:
                sec_check_cache[key] = validate_section(sec_df, county, town, sec, sub)
            sec_ok, sec_note, sec_low = sec_check_cache[key]
            if sec_note:
                notes.append(sec_note)
            if sec_low:
                conf_low = True

            if p["mother"] is None:
                display = ""
                conf_low = True
            else:
                display = landno_display(p["mother"], p["child"])
                if display != p["raw"] and f"由範圍「" not in "".join(p["notes"]):
                    # 正規化後顯示值與原文不同（如 0105→105、1146-0002→1146-2）不需備註，
                    # 但保留原文欄讓使用者可比對
                    pass

            out_rows.append([
                f"{seq:03d}",
                _clean_cell(site["site_name"]),
                normalize_county(county),
                town,
                re.sub(r"段$", "", sec) if sec else "",
                re.sub(r"小段$", "", sub) if sub else "",
                p["raw"],
                p["mother"] or "",
                p["child"] or "",
                p["annotation"],
                _clean_cell(site["zone_in"]),
                _clean_cell(site["landuse_in"]),
                _clean_cell(site["fz_in"]),
                "低" if conf_low else "高",
                "；".join(dict.fromkeys(n for n in notes if n)),
                _clean_cell(site["agency"]),
                _clean_cell(site["address"]),
            ])

    # ---- 寫 xlsx ----
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "中繼檔"
    yellow = PatternFill(start_color="FFF2A0", end_color="FFF2A0", fill_type="solid")
    bold = Font(bold=True)

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = bold
    for row in out_rows:
        ws.append(row)
        if row[13] == "低":
            for cell in ws[ws.max_row]:
                cell.fill = yellow

    widths = [7, 20, 8, 10, 12, 10, 14, 8, 8, 10, 16, 16, 46, 8, 60, 16, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)

    # ---- 摘要 ----
    n_low = sum(1 for r in out_rows if r[13] == "低")
    n_sites = len({r[0] for r in out_rows})
    print(f"已寫入 {OUT_XLSX}")
    print(f"案場 {n_sites}、地號列 {len(out_rows)}、低信心列 {n_low}")
    if issues:
        print("\n--- 需注意 ---")
        for msg in issues:
            print(msg)

    print("\n--- 低信心案場 ---")
    seen = set()
    for r in out_rows:
        if r[13] == "低" and r[0] not in seen:
            seen.add(r[0])
            print(f"[{r[0]}] {r[1]}：{r[14][:80]}")


if __name__ == "__main__":
    main()
