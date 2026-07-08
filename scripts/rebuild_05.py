# -*- coding: utf-8 -*-
"""從 02_比對結果 + 出圖成果資料夾，重建涵蓋全部案場的 05 出圖清單。

（單案/分批出圖會覆蓋 make_maps 當次的 05；此腳本掃描實際成果統整全部。）
系統 Python 執行。
"""
import re
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

REPO = Path(__file__).resolve().parent.parent
CASE = "20260707-清潔隊部資訊"
OUT = REPO / "outputs" / CASE
RESULT = OUT / "出圖成果"
THEMES = ["使用分區", "使用地編定", "國土功能分區", "正射套疊"]


def parcels_text(sub: pd.DataFrame) -> str:
    groups = defaultdict(list)
    for _, r in sub.iterrows():
        seg = str(r.get("段") or "")
        if seg and not seg.endswith("段"):
            seg += "段"
        ss = str(r.get("小段") or "")
        if ss and ss not in ("", "None", "nan"):
            seg += ss + "小段"
        groups[seg].append(str(r.get("地號原文") or ""))
    return "；".join(f"{s} {'、'.join(n)}" for s, n in groups.items())


def main():
    df = pd.read_excel(OUT / "02_比對結果.xlsx", dtype={"案場ID": str})
    df = df[df["比對狀態"] == "成功"]
    meta = {}
    for cid, sub in df.groupby("案場ID"):
        cid3 = str(cid).strip().zfill(3)
        meta[cid3] = {
            "名稱": str(sub.iloc[0]["案場名稱"]),
            "縣市": str(sub.iloc[0]["縣市"]),
            "鄉鎮區": str(sub.iloc[0].get("鄉鎮區") or ""),
            "地號": parcels_text(sub),
        }

    rows = []
    for folder in sorted(RESULT.iterdir()):
        if not folder.is_dir():
            continue
        m = re.match(r"^(\d+)_(.+)$", folder.name)
        if not m:
            continue
        cid = m.group(1)
        info = meta.get(cid, {"名稱": m.group(2), "縣市": "", "鄉鎮區": "", "地號": ""})
        for theme in THEMES:
            png = folder / f"{theme}.png"
            rows.append([cid, info["名稱"], info["縣市"], info["鄉鎮區"], theme,
                         "成功" if png.exists() else "缺圖",
                         f"出圖成果/{folder.name}/{theme}.png", info["地號"]])

    wb = Workbook()
    ws = wb.active
    ws.title = "出圖清單"
    header = ["案場ID", "案場名稱", "縣市", "鄉鎮區", "主題", "狀態", "檔案路徑", "涵蓋地號"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate([7, 20, 8, 8, 12, 6, 45, 55], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    ws.freeze_panes = "A2"
    wb.save(OUT / "05_出圖清單.xlsx")
    n_case = len({r[0] for r in rows})
    n_miss = sum(1 for r in rows if r[5] == "缺圖")
    print(f"已重建 05_出圖清單.xlsx：{n_case} 案場、{len(rows)} 列、缺圖 {n_miss}")


if __name__ == "__main__":
    main()
