# -*- coding: utf-8 -*-
"""fallback 鏈寫檔端:把 04_新舊地號對照 回填進 01_中繼檔(spec §5.2)。

- 查詢狀態=成功:該列改用查得的段/小段/地號(宗地母/子重算),原值記備註。
- 查詢狀態=查無:該列標「比對模式=母號範圍」(match_parcels 會撈同段同宗地母全部宗地),
  地號原文改為母號,原值記備註。
- 查得新段名/查得新地號含「、」(query_landno 多候選串接):視為 ambiguous,不回填、
  計入 skip、備註不動,main() 會列出警告。
- 04 裡對不上 01 任何列的 (案場ID,原地號) 會回報在 stats["unmatched_lookup"],
  main() 會逐筆印警告。
- main() 就地編輯 01_中繼檔.xlsx(load_workbook 非 read_only),只改動有變動的
  儲存格,保留既有格式(黃底等)與分頁名稱;若原檔無「比對模式」欄,在表頭尾端補一欄。
用法:python -X utf8 scripts/apply_landno.py <案子名>
"""
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
COLS_EXTRA = "比對模式"


def split_section(name: str) -> tuple[str, str]:
    """'聚興段新興小段'→('聚興','新興');'後壁段'→('後壁','');'民權一段'→('民權一','')。"""
    name = (name or "").strip()
    m = re.match(r"^(.*?)段(?:(.*?)小段)?$", name)
    if not m:
        return name, ""
    return m.group(1), (m.group(2) or "")


def _parse_no(s: str) -> tuple[str, str]:
    s = str(s).strip().split(".")[0]
    m = re.match(r"^(\d+)(?:-(\d+))?$", s)
    if not m:
        return "", ""
    return f"{int(m.group(1)):04d}", (f"{int(m.group(2)):04d}" if m.group(2) else "0000")


def apply_corrections(meta_rows: list[dict], lookup_rows: list[dict]):
    """依 04 修正 01 列。key=(案場ID, 原地號)。回傳 (新列list, 統計dict)。

    stats: {"exact": n, "mother": n, "skip": n, "unmatched_lookup": [(案場ID, 原地號), ...]}
    unmatched_lookup 收「04 有、但 01 對不上」的列;skip 含 mother 分支失敗與多值
    ambiguous 兩種情況。
    """
    idx = {}
    for lr in lookup_rows:
        idx[(str(lr.get("案場ID", "")).strip().zfill(3),
             str(lr.get("原地號", "")).strip())] = lr
    stats = {"exact": 0, "mother": 0, "skip": 0, "unmatched_lookup": []}
    matched_keys = set()
    out = []
    for row in meta_rows:
        r = dict(row)
        key = (str(r.get("案場ID", "")).strip().zfill(3),
               str(r.get("地號原文", "")).strip())
        lr = idx.get(key)
        if lr is None:
            out.append(r)
            continue
        matched_keys.add(key)
        orig = f"{r['段']}{('段' if r['段'] else '')}{(str(r['小段']) + '小段') if str(r.get('小段') or '').strip() not in ('', 'None', 'nan') else ''}{r['地號原文']}"
        new_seg_raw = str(lr.get("查得新段名") or "")
        new_no_raw = str(lr.get("查得新地號") or "")
        if "、" in new_seg_raw or "、" in new_no_raw:
            # 多候選串接,無法判斷取哪一個 → 不回填、視為 skip、備註不動
            stats["skip"] += 1
            out.append(r)
            continue
        if str(lr.get("查詢狀態", "")).strip() == "成功":
            seg, sub = split_section(new_seg_raw)
            mother, child = _parse_no(new_no_raw)
            if not seg or not mother:
                stats["skip"] += 1
                out.append(r)
                continue
            r["段"], r["小段"], r["地號原文"] = seg, sub, new_no_raw.strip()
            r["宗地母"], r["宗地子"] = mother, child
            r["備註"] = (str(r.get("備註") or "") +
                         f"｜原{orig},依新舊地號對照回填").lstrip("｜")
            stats["exact"] += 1
        else:  # 查無 → 母號範圍
            mother = str(r.get("宗地母", "")).strip().split(".")[0]
            if not mother:
                stats["skip"] += 1
                out.append(r)
                continue
            r[COLS_EXTRA] = "母號範圍"
            r["地號原文"] = str(int(mother))
            r["宗地母"], r["宗地子"] = mother.zfill(4), ""
            r["備註"] = (str(r.get("備註") or "") +
                         f"｜原{orig}查無對應,以母號範圍示意出圖").lstrip("｜")
            stats["mother"] += 1
        out.append(r)
    for key, lr in idx.items():
        if key not in matched_keys:
            stats["unmatched_lookup"].append(
                (str(lr.get("案場ID", "")).strip(), str(lr.get("原地號", "")).strip()))
    return out, stats


CHANGE_FIELDS = ["段", "小段", "地號原文", "宗地母", "宗地子", "備註", COLS_EXTRA]


def _read_lookup(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c) for c in rows[0]]
    return [dict(zip(hdr, [("" if v is None else v) for v in row]))
            for row in rows[1:] if any(v is not None for v in row)]


def main():
    from openpyxl import load_workbook
    case = sys.argv[1]
    d = REPO / "outputs" / case
    meta_path = d / "01_中繼檔.xlsx"

    wb = load_workbook(meta_path)  # 非 read_only:保留格式,可就地編輯
    ws = wb.active

    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    while hdr and hdr[-1] is None:
        hdr.pop()
    hdr = [str(h) for h in hdr]
    col_of = {h: i + 1 for i, h in enumerate(hdr)}

    meta_rows, row_nums = [], []
    for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None for v in row):
            continue
        vals = list(row[:len(hdr)]) + [None] * max(0, len(hdr) - len(row))
        meta_rows.append(dict(zip(hdr, [("" if v is None else v) for v in vals])))
        row_nums.append(r_idx)

    lookup = _read_lookup(d / "04_新舊地號對照清單.xlsx")
    out, stats = apply_corrections(meta_rows, lookup)

    if COLS_EXTRA not in col_of:
        new_col = len(hdr) + 1
        ws.cell(row=1, column=new_col, value=COLS_EXTRA)
        col_of[COLS_EXTRA] = new_col
        hdr.append(COLS_EXTRA)

    for orig_row, r, r_idx in zip(meta_rows, out, row_nums):
        if r == orig_row:
            continue
        for h in CHANGE_FIELDS:
            new_v, old_v = r.get(h, ""), orig_row.get(h, "")
            if new_v != old_v:
                ws.cell(row=r_idx, column=col_of[h], value=new_v)

    wb.save(meta_path)
    print(f"回填完成:精確 {stats['exact']}、母號範圍 {stats['mother']}、略過 {stats['skip']}")

    if stats["unmatched_lookup"]:
        print(f"警告:{len(stats['unmatched_lookup'])} 筆04對照在01中找不到對應列(案場ID/原地號):")
        for cid, landno in stats["unmatched_lookup"]:
            print(f"  - {cid} / {landno}")

    ambiguous = [lr for lr in lookup
                 if "、" in str(lr.get("查得新段名") or "") or "、" in str(lr.get("查得新地號") or "")]
    if ambiguous:
        print(f"警告:{len(ambiguous)} 筆04對照為多值(ambiguous),未回填、已計入略過:")
        for lr in ambiguous:
            print(f"  - 案場{lr.get('案場ID')} 原地號{lr.get('原地號')}: "
                  f"新段名={lr.get('查得新段名')} 新地號={lr.get('查得新地號')}")


if __name__ == "__main__":
    main()
