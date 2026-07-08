# -*- coding: utf-8 -*-
"""把 05 出圖清單 CSV 轉為格式化 xlsx。

**獨立執行、不 import qgis/gdal**：openpyxl 依賴 lxml，與 GDAL 共用 libxml2；
若在跑過大量 GDAL 圖磚處理的同一進程建立 openpyxl Workbook 會存取違例硬崩潰
（Windows fatal exception: access violation）。故 xlsx 一律由本乾淨進程產生。

用法：python csv_to_xlsx.py <in.csv> <out.xlsx>
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

COL_WIDTHS = [7, 18, 12, 40, 9, 50, 17, 8, 40]


def main() -> None:
    if len(sys.argv) < 3:
        sys.exit("用法：python csv_to_xlsx.py <in.csv> <out.xlsx>")
    csv_path, xlsx_path = sys.argv[1], sys.argv[2]
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        sys.exit(f"CSV 無內容：{csv_path}")

    wb = Workbook()
    ws = wb.active
    ws.title = "出圖清單"
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, w in enumerate(COL_WIDTHS, start=1):
        if i <= ws.max_column:
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    print(f"已產生 {xlsx_path}（{len(rows) - 1} 列）")


if __name__ == "__main__":
    main()
