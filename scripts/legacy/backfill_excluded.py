# -*- coding: utf-8 -*-
"""把 5 個排除案場（新舊地號對照結果）回填進 02_比對結果.xlsx，供 make_maps 出圖。

依 04_新舊地號對照清單：
- 021/038/086：查得精確對照地號（存在於 111年2月地籍）→ 精確出圖
- 023/063：新舊地號系統查無 → 以「母號範圍」出圖（同段、同宗地母之所有子號），圖面註記

FID 慣例與 match_parcels 一致：pyogrio 讀全表 enumerate 的 0-based 序位。
在系統 Python 執行（pyogrio+openpyxl 同進程安全；libxml2 衝突僅見於 QGIS python 大量圖磚後）。
"""
import warnings
from pathlib import Path

import pyogrio
from openpyxl import load_workbook

warnings.filterwarnings("ignore")
REPO = Path(__file__).resolve().parent.parent
CAD = REPO / "地籍圖資"
OUT02 = REPO / "outputs" / "20260707-清潔隊部資訊" / "02_比對結果.xlsx"

# (案場ID, 案場名稱, 縣市, 鄉鎮區, shp檔名, 段, 小段, 母號list, 模式, 地號原文, 輸入分區, 輸入地目, 備註)
#   模式 exact = 子號0000 精確；all = 該母號所有子號（母號範圍）
TARGETS = [
    (21, "蘆竹區中隊", "桃園市", "蘆竹區", "桃園市", "後壁", "", [12], "exact",
     "12（原坑子口段後壁厝小段1-1，重測後新段地號）", "森林區", "林業用地",
     "原地號為重測前舊地號；依04對照回填重測後後壁段12（已驗證存在111年2月地籍）"),
    (38, "潭子區清潔隊", "臺中市", "潭子區", "台中市", "聚興", "新興", [461], "exact",
     "461（原寶興段176，對照重測前舊段地號）", "特定農業區", "農牧用地、水利用地、交通用地",
     "原寶興段為重測後段名；111年2月地籍以舊段名聚興段新興小段收錄，依04對照回填461。"
     "案場其餘地號(178~201)尚未個別對照，本圖僅呈現176之對照地號"),
    (86, "高樹鄉清潔隊", "屏東縣", "高樹鄉", "屏東縣", "東振新", "", [707], "exact",
     "707（原東振段1215，對照重測前舊段地號）", "一般農業區", "墳墓用地",
     "依04對照回填東振新段707（已驗證存在111年2月地籍）"),
    (23, "觀音區中隊", "桃園市", "觀音區", "桃園市", "草漯", "",
     [2529, 2530, 2531, 2532, 2533, 2535], "all",
     "母號範圍", "森林區", "國土安全用地",
     "原2535-1新舊地號系統查無（疑逕為分割）；以母號範圍(草漯段2529~2533、2535之所有宗地)示意出圖"),
    (63, "仁德區清潔隊", "臺南市", "仁德區", "台南市", "崑崙", "", [786], "all",
     "母號範圍", "一般農業區", "交通用地",
     "原786-4新舊地號系統查無（疑逕為分割）；以母號範圍(崑崙段786之所有宗地)示意出圖"),
]

# 02 欄位順序
COLS = ["案場ID", "案場名稱", "縣市", "鄉鎮區", "段", "小段", "地號原文", "宗地母", "宗地子",
        "部分註記", "輸入使用分區", "輸入使用地目", "輸入國土功能分區", "比對狀態", "小段碼",
        "圖資FID", "圖資使分", "圖資使類", "登記面積m2", "圖資國土功能分區(面積佔比)",
        "使分一致", "使類一致", "功能分區一致", "備註"]

# 依縣市快取讀取的地籍屬性
_cache = {}


def load_cad(shp_name):
    if shp_name not in _cache:
        df = pyogrio.read_dataframe(str(CAD / f"{shp_name}.shp"), read_geometry=False,
                                    columns=["段", "小段", "宗地母", "宗地子", "小段碼",
                                             "使分", "使類", "登面積"])
        _cache[shp_name] = df
    return _cache[shp_name]


def build_rows():
    new_rows = []
    for (cid, name, county, town, shp, seg, ssub, mothers, mode,
         landno_disp, in_zone, in_type, note) in TARGETS:
        df = load_cad(shp)
        seg_s = df["段"].astype(str)
        ss_s = df["小段"].astype(str)
        m_s = df["宗地母"].astype(str)
        c_s = df["宗地子"].astype(str)
        mset = {f"{m:04d}" for m in mothers}
        for m in mothers:
            m4 = f"{m:04d}"
            mask = (seg_s == seg) & (m_s == m4)
            if ssub:
                mask &= (ss_s == ssub)
            if mode == "exact":
                mask &= (c_s == "0000")
            fids = [i for i, v in enumerate(mask) if v]
            if not fids:
                print(f"[警告] 案場{cid} {seg}{ssub} 母{m4} 模式{mode}：查無宗地")
                continue
            hit = df.iloc[fids]
            zones = sorted({str(v) for v in hit["使分"].dropna() if str(v) != "None"})
            types = sorted({str(v) for v in hit["使類"].dropna() if str(v) != "None"})
            codes = sorted({str(v) for v in hit["小段碼"].dropna()})
            area = float(hit["登面積"].astype(float).sum())
            disp = landno_disp if mode == "exact" else f"{m}（母號範圍）"
            row = {c: "" for c in COLS}
            row.update({
                "案場ID": f"{cid:03d}",  # 與現有 02/檔名一致的零填字串
                "案場名稱": name, "縣市": county, "鄉鎮區": town,
                "段": seg, "小段": ssub, "地號原文": disp,
                "宗地母": m4, "宗地子": ("0000" if mode == "exact" else ""),
                "輸入使用分區": in_zone, "輸入使用地目": in_type, "輸入國土功能分區": "",
                "比對狀態": "成功", "小段碼": "；".join(codes),
                "圖資FID": ";".join(str(f) for f in fids),
                "圖資使分": "、".join(zones), "圖資使類": "、".join(types),
                "登記面積m2": area, "備註": note,
            })
            new_rows.append(row)
            print(f"案場{cid} {seg}{ssub} 母{m4}（{mode}）：{len(fids)} 宗地，"
                  f"使分={zones}，使類={types}")
    return new_rows


def main():
    rows = build_rows()
    wb = load_workbook(OUT02)
    ws = wb.active
    header = [c.value for c in ws[1]]
    # 先移除舊的這些案場ID回填列（若重跑）；相容 int 與零填字串兩種寫法
    idset = {f"{t[0]:03d}" for t in TARGETS}
    st_col = header.index("比對狀態") + 1
    to_del = [r for r in range(ws.max_row, 1, -1)
              if str(ws.cell(r, 1).value or "").strip().zfill(3) in idset
              and str(ws.cell(r, st_col).value) == "成功"]
    for r in to_del:
        ws.delete_rows(r, 1)
    for row in rows:
        ws.append([row.get(h, "") for h in header])
    wb.save(OUT02)
    print(f"\n已回填 {len(rows)} 列到 {OUT02}（刪除舊回填 {len(to_del)} 列）")


if __name__ == "__main__":
    main()
