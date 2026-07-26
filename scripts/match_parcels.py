# -*- coding: utf-8 -*-
"""
階段2比對：01_中繼檔.xlsx → 02_比對結果.xlsx ＋ 03_無法對應清單.xlsx

流程（設計文件 §5.2）：
1. 段名 → 小段碼：查 config/section_index.csv（現行地籍段名索引；鄉鎮寬鬆比對）
2. 依（小段碼, 宗地母, 宗地子）從縣市地籍 shp 篩宗地（屬性先查、再按 FID 讀幾何）
3. 屬性取值：使分/使類/登面積；國土功能分區＝宗地 polygon 與縣市功能分區 shp
   空間交集（面積佔比 ≥1% 列出）
4. 02＝全部地號列（含輸入值 vs 圖資值一致性檢核）；03＝失敗列
   （查無段小段／查無地號／縣市檔缺漏／其他），供階段2b新舊地號查詢

執行：python -X utf8 scripts/match_parcels.py [案子名]
     （預設 20260707-清潔隊部資訊）
"""

from __future__ import annotations

import re
import sys
import warnings
from collections import defaultdict
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))

import geopandas as gpd
import pandas as pd
import pyogrio
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from parse_helpers import normalize_county, tai_norm, town_matches
from zone_utils import extract_zones_from_text, normalize_zone

CASE_NAME = sys.argv[1] if len(sys.argv) > 1 else "20260707-清潔隊部資訊"
OUT_DIR = REPO_ROOT / "outputs" / CASE_NAME
IN_XLSX = OUT_DIR / "01_中繼檔.xlsx"
OUT_02 = OUT_DIR / "02_比對結果.xlsx"
OUT_03 = OUT_DIR / "03_無法對應清單.xlsx"

SECTION_INDEX_CSV = REPO_ROOT / "config" / "section_index.csv"
REGISTRY_YAML = REPO_ROOT / "config" / "county_registry.yaml"

# 功能分區佔比列出門檻（設計文件 §5.2：<1% 忽略，可調）
ZONE_SHARE_MIN = 0.01
# 一致性檢核採用門檻（低於此佔比的分區不參與「不一致」判定，避免邊界切割雜訊）
ZONE_SHARE_PRIMARY = 0.05

# 有「總圖/合併檔」的縣市只用合併檔，避免與分類檔重複計算
PREFERRED_ZONE_FILES = {
    "苗栗縣": ["國土功能分區1091208修"],
    "臺中市": ["功能分區總圖"],
    "臺東縣": ["臺東陸域功能分區D", "臺東海域功能分區D"],
}


# ---------------------------------------------------------------------------
# 段名 → 小段碼
# ---------------------------------------------------------------------------


def load_section_index() -> pd.DataFrame:
    df = pd.read_csv(SECTION_INDEX_CSV, dtype=str, encoding="utf-8-sig")
    for col in ("縣市", "鄉鎮", "段", "小段"):
        df[f"{col}_t"] = df[col].map(tai_norm)
    return df


def resolve_subsection_codes(sec_df, county, town, sec, sub) -> list[str]:
    """回傳候選小段碼清單（一般恰 1 筆；0 筆=查無；>1 筆=同名段須靠地號分辨）。"""
    county_n = normalize_county(county)
    mask = (sec_df["縣市_t"] == county_n) & (sec_df["段_t"] == tai_norm(sec))
    if sub and not pd.isna(sub):
        mask &= sec_df["小段_t"] == tai_norm(sub)
    else:
        mask &= sec_df["小段"].isna()
    cand = sec_df.loc[mask]
    if cand.empty:
        return []
    town_hit = cand[cand["鄉鎮"].map(lambda t: town_matches(t, tai_norm(town)))]
    picked = town_hit if not town_hit.empty else cand
    return list(dict.fromkeys(picked["小段碼"].dropna()))


# ---------------------------------------------------------------------------
# 地籍宗地比對
# ---------------------------------------------------------------------------

CADASTRAL_ATTR_COLS = ["小段碼", "宗地母", "宗地子", "使分", "使類", "登面積"]


def pick_fids(key_to_fids: dict, key_mother: dict, code: str,
              mother: str, child: str, mode: str) -> list[int]:
    """依比對模式取 fids：exact=(code,母,子)；母號範圍=(code,母)全部子號。"""
    if mode == "母號範圍":
        return list(key_mother.get((code, mother), []))
    return list(key_to_fids.get((code, mother, child), []))


def match_county_parcels(shp_path: Path, wanted: list[dict]) -> None:
    """就地更新 wanted 各列：找到→填 fids/使分/使類/登面積；找不到→維持未命中。

    wanted 每筆需有 candidates(小段碼清單)/宗地母/宗地子 欄位。
    """
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=RuntimeWarning)
        attrs = pyogrio.read_dataframe(
            str(shp_path), read_geometry=False, columns=CADASTRAL_ATTR_COLS
        )
    key_to_fids: dict[tuple, list[int]] = defaultdict(list)
    key_mother: dict[tuple, list[int]] = defaultdict(list)
    for fid, (code, m, c) in enumerate(
        zip(attrs["小段碼"], attrs["宗地母"], attrs["宗地子"])
    ):
        key_to_fids[(code, m, c)].append(fid)
        key_mother[(code, m)].append(fid)

    for w in wanted:
        for code in w["candidates"]:
            fids = pick_fids(key_to_fids, key_mother, code,
                             w["宗地母"], w.get("宗地子", ""), w.get("mode", ""))
            if fids:
                rows = attrs.iloc[fids]
                w["fids"] = fids
                w["小段碼"] = code
                w["使分"] = "、".join(sorted(set(str(v) for v in rows["使分"].dropna()))) or None
                w["使類"] = "、".join(sorted(set(str(v) for v in rows["使類"].dropna()))) or None
                w["登面積"] = float(rows["登面積"].astype(float).sum())
                break


def read_parcel_geometries(shp_path: Path, fids: list[int]) -> gpd.GeoDataFrame:
    if not fids:
        # 某縣市整批比對不到時 fids 為空:仍要帶 fid 欄,否則呼叫端 geoms["fid"] KeyError
        return gpd.GeoDataFrame({"fid": []}, geometry=[], crs="EPSG:3826")
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=RuntimeWarning)
        gdf = pyogrio.read_dataframe(str(shp_path), fids=sorted(set(fids)), columns=[])
    gdf["fid"] = sorted(set(fids))
    if gdf.crs is None:
        gdf = gdf.set_crs("EPSG:3826")
    return gdf


# ---------------------------------------------------------------------------
# 功能分區圖層載入與交集
# ---------------------------------------------------------------------------


def _fix_crs(gdf: gpd.GeoDataFrame) -> gpd.GeoDataFrame:
    """依 registry 探勘結論修 CRS：缺 .prj／自訂 TM 參數＝EPSG:3826；3825→轉 3826。"""
    if gdf.crs is None:
        return gdf.set_crs("EPSG:3826")
    epsg = gdf.crs.to_epsg()
    if epsg == 3826:
        return gdf
    if epsg == 3825:
        return gdf.to_crs("EPSG:3826")
    if epsg is None:
        # 自訂 PROJCS（中央經線121、比例0.9999、假東250000）實為 TWD97 TM2
        return gdf.set_crs("EPSG:3826", allow_override=True)
    return gdf.to_crs("EPSG:3826")


def load_county_zones(registry: dict, county: str) -> gpd.GeoDataFrame | None:
    """載入縣市功能分區為單一 GeoDataFrame（欄位 zone=標準分區名）。無圖資回 None。"""
    entries = registry["counties"].get(county, {}).get("functional_zones", [])
    if not entries:
        return None
    preferred = PREFERRED_ZONE_FILES.get(county)
    if preferred:
        entries = [e for e in entries if any(p in e["path"] for p in preferred)]

    frames = []
    for e in entries:
        path = REPO_ROOT / e["path"]
        stem_label = normalize_zone(Path(e["path"]).stem)
        zone_field = e.get("zone_field")
        if zone_field is None and stem_label is None:
            continue  # 非四大分區圖層（行政區界、未來發展、纜線…）
        try:
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", category=RuntimeWarning)
                cols = [zone_field] if zone_field else []
                gdf = gpd.GeoDataFrame(
                    pyogrio.read_dataframe(str(path), columns=cols)
                )
        except Exception as exc:  # noqa: BLE001
            print(f"  [警告] {county} 功能分區讀取失敗：{e['path']}（{exc}）")
            continue
        gdf = _fix_crs(gdf)
        if zone_field:
            gdf["zone"] = gdf[zone_field].map(lambda v: normalize_zone(v) or stem_label)
        else:
            gdf["zone"] = stem_label
        gdf = gdf[gdf["zone"].notna()][["zone", "geometry"]]
        if not gdf.empty:
            frames.append(gdf)

    if not frames:
        return None
    zones = gpd.GeoDataFrame(pd.concat(frames, ignore_index=True), crs="EPSG:3826")
    return zones


def intersect_zones(geom, zones: gpd.GeoDataFrame) -> list[tuple[str, float]]:
    """回傳 [(標準分區名, 面積佔比), ...]，佔比 ≥ ZONE_SHARE_MIN，由大到小。"""
    if geom is None or geom.is_empty or zones is None:
        return []
    idx = zones.sindex.query(geom, predicate="intersects")
    if len(idx) == 0:
        return []
    shares: dict[str, float] = defaultdict(float)
    total = geom.area
    if total <= 0:
        return []
    for _, row in zones.iloc[idx].iterrows():
        try:
            inter = geom.intersection(row.geometry)
        except Exception:  # noqa: BLE001 — 個別無效幾何以 buffer(0) 修復重試
            inter = geom.buffer(0).intersection(row.geometry.buffer(0))
        if not inter.is_empty:
            shares[row["zone"]] += inter.area
    result = [(z, a / total) for z, a in shares.items() if a / total >= ZONE_SHARE_MIN]
    return sorted(result, key=lambda t: -t[1])


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

OUT02_COLUMNS = [
    "案場ID", "案場名稱", "縣市", "鄉鎮區", "段", "小段", "地號原文", "宗地母", "宗地子",
    "部分註記", "輸入使用分區", "輸入使用地目", "輸入國土功能分區",
    "比對狀態", "小段碼", "圖資FID", "圖資使分", "圖資使類", "登記面積m2",
    "圖資國土功能分區(面積佔比)", "使分一致", "使類一致", "功能分區一致", "備註",
]

OUT03_COLUMNS = [
    "案場ID", "案場名稱", "縣市", "鄉鎮區", "段", "小段", "地號原文", "宗地母", "宗地子",
    "失敗分類", "說明", "備註",
]


# 新舊名稱同義詞（比對視為一致）：非都市土地使用地類別「墳墓用地」現多寫「殯葬用地」
_SYNONYMS = {"殯葬用地": "墳墓用地"}


def _syn(s: str) -> str:
    for old, new in _SYNONYMS.items():
        s = s.replace(old, new)
    return s


def consistency(gis_value, input_text) -> str:
    if gis_value is None or gis_value == "":
        return "圖資無值"
    if not isinstance(input_text, str) or not input_text.strip():
        return "無輸入值"
    input_n = _syn(input_text)
    return "一致" if all(_syn(part) in input_n for part in gis_value.split("、")) else "不一致"


def zone_consistency(zone_shares, input_text, has_zone_data: bool) -> str:
    if not has_zone_data:
        return "無功能分區圖資"
    input_zones = extract_zones_from_text(input_text)
    primary = [z for z, s in zone_shares if s >= ZONE_SHARE_PRIMARY]
    if not primary:
        return "圖資無交集"
    if not input_zones:
        return "無法比對(輸入無分區值)"
    return "一致" if all(z in input_zones for z in primary) else "不一致"


def write_xlsx(path: Path, columns: list[str], rows: list[list], highlight_col: str,
               highlight_values: set[str], col_widths: list[int]) -> None:
    wb = Workbook()
    ws = wb.active
    yellow = PatternFill(start_color="FFF2A0", end_color="FFF2A0", fill_type="solid")
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    hi = columns.index(highlight_col) if highlight_col in columns else None
    for row in rows:
        ws.append(row)
        if hi is not None and str(row[hi]) in highlight_values:
            for cell in ws[ws.max_row]:
                cell.fill = yellow
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = wrap
    ws.freeze_panes = "A2"
    wb.save(path)


def main() -> None:
    df = pd.read_excel(IN_XLSX, dtype=str)
    print(f"讀入中繼檔 {len(df)} 列")

    sec_df = load_section_index()
    with REGISTRY_YAML.open(encoding="utf-8") as f:
        registry = yaml.safe_load(f)

    # -- 每列建 work dict --
    works: list[dict] = []
    for _, r in df.iterrows():
        w = dict(r)
        w["fids"] = None
        w["小段碼"] = None
        w["使分"] = w["使類"] = w["登面積"] = None
        w["status"] = None
        w["zone_shares"] = []
        raw_mode = r.get("比對模式")
        w["mode"] = "" if pd.isna(raw_mode) else str(raw_mode).strip()
        raw_child = w.get("宗地子")
        w["宗地子"] = "" if pd.isna(raw_child) else str(raw_child).strip()
        county = r["縣市"]
        if pd.isna(w.get("宗地母")) or not str(w.get("宗地母")).strip():
            w["status"] = "其他"
            w["status_note"] = "中繼檔宗地母為空（解析失敗列）"
            w["candidates"] = []
        elif not registry["counties"].get(county, {}).get("cadastral"):
            w["status"] = "縣市檔缺漏"
            w["status_note"] = f"{county} 無地籍圖資"
            w["candidates"] = []
        else:
            cands = resolve_subsection_codes(
                sec_df, county, r["鄉鎮區"], r["段"], r.get("小段")
            )
            w["candidates"] = cands
            if not cands:
                w["status"] = "查無段小段"
                w["status_note"] = "地籍索引查無此段/小段"
        works.append(w)

    # -- 依縣市比對宗地 --
    by_county: dict[str, list[dict]] = defaultdict(list)
    for w in works:
        if w["status"] is None:
            by_county[w["縣市"]].append(w)

    geom_by_key: dict[tuple, object] = {}
    zones_cache: dict[str, gpd.GeoDataFrame | None] = {}

    for county, items in by_county.items():
        shp_path = REPO_ROOT / registry["counties"][county]["cadastral"]
        print(f"{county}：比對 {len(items)} 筆地號…", flush=True)
        match_county_parcels(shp_path, items)

        all_fids = sorted({f for w in items if w["fids"] for f in w["fids"]})
        geoms = read_parcel_geometries(shp_path, all_fids)
        fid_geom = dict(zip(geoms["fid"], geoms.geometry))

        zones_cache[county] = load_county_zones(registry, county)
        zones = zones_cache[county]

        for w in items:
            if not w["fids"]:
                w["status"] = "查無地號"
                w["status_note"] = "段/小段存在，但查無此宗地母/子（可能為地籍異動）"
                continue
            w["status"] = "成功"
            w["status_note"] = ""
            geom_list = [fid_geom[f] for f in w["fids"] if f in fid_geom]
            if len(geom_list) == 1:
                geom = geom_list[0]
            else:
                from shapely.ops import unary_union

                geom = unary_union(geom_list)
                w["status_note"] = f"地籍有 {len(geom_list)} 筆同地號，幾何已合併"
            geom_by_key[(w["縣市"], w["小段碼"], w["宗地母"], w["宗地子"])] = geom
            w["zone_shares"] = intersect_zones(geom, zones)

    # -- 組 02 / 03 --
    rows02: list[list] = []
    rows03: list[list] = []
    for w in works:
        status = w["status"]
        zone_txt = "、".join(f"{z}({s:.0%})" for z, s in w["zone_shares"])
        has_zone = zones_cache.get(w["縣市"]) is not None
        note = "；".join(x for x in [w.get("status_note", ""), str(w.get("備註") or "")] if x and x != "nan")

        rows02.append([
            w["案場ID"], w["案場名稱"], w["縣市"], w["鄉鎮區"],
            w.get("段"), w.get("小段"), w["地號原文"], w.get("宗地母"), w.get("宗地子"),
            w.get("部分註記"), w.get("輸入使用分區"), w.get("輸入使用地目"),
            w.get("輸入國土功能分區"),
            status, w.get("小段碼"),
            ";".join(str(f) for f in (w["fids"] or [])),
            w.get("使分"), w.get("使類"),
            round(w["登面積"], 2) if w.get("登面積") is not None else None,
            zone_txt,
            consistency(w.get("使分"), w.get("輸入使用分區")) if status == "成功" else "",
            consistency(w.get("使類"), w.get("輸入使用地目")) if status == "成功" else "",
            zone_consistency(w["zone_shares"], w.get("輸入國土功能分區"), has_zone)
            if status == "成功" else "",
            note,
        ])

        if status != "成功":
            rows03.append([
                w["案場ID"], w["案場名稱"], w["縣市"], w["鄉鎮區"],
                w.get("段"), w.get("小段"), w["地號原文"], w.get("宗地母"), w.get("宗地子"),
                status, w.get("status_note", ""), str(w.get("備註") or "").replace("nan", ""),
            ])

    write_xlsx(OUT_02, OUT02_COLUMNS, rows02, "比對狀態",
               {"查無段小段", "查無地號", "縣市檔缺漏", "其他"},
               [7, 18, 8, 9, 10, 8, 12, 7, 7, 8, 14, 14, 40, 10, 10, 12, 12, 12, 12, 42, 9, 9, 16, 40])
    write_xlsx(OUT_03, OUT03_COLUMNS, rows03, "失敗分類",
               {"查無段小段", "查無地號", "縣市檔缺漏", "其他"},
               [7, 18, 8, 9, 10, 8, 12, 7, 7, 12, 40, 40])

    # -- 摘要 --
    n_ok = sum(1 for w in works if w["status"] == "成功")
    print(f"\n已寫入 {OUT_02}")
    print(f"已寫入 {OUT_03}")
    print(f"成功 {n_ok}/{len(works)}；失敗 {len(rows03)} 筆")
    fail_stat = defaultdict(int)
    for w in works:
        if w["status"] != "成功":
            fail_stat[w["status"]] += 1
    for k, v in fail_stat.items():
        print(f"  {k}: {v}")

    incons = [r for r in rows02 if "不一致" in (str(r[20]), str(r[21]), str(r[22]))]
    print(f"輸入值與圖資值不一致：{len(incons)} 筆（詳見 02 檔）")


if __name__ == "__main__":
    main()
