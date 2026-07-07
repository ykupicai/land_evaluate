# -*- coding: utf-8 -*-
"""階段2b：新舊地號自動查詢（03_無法對應清單 → 04_新舊地號對照清單）。

用法（系統 Python）：
    python -X utf8 scripts/query_landno.py [案子名]
    （預設案子：20260707-清潔隊部資訊）

讀 outputs/<案子>/03_無法對應清單.xlsx，逐筆到各縣市「地籍圖重測地段新舊地建號
查詢」系統查對照，寫出 outputs/<案子>/04_新舊地號對照清單.xlsx。

==============================================================================
各縣市系統探勘紀錄（2026-07-07 實測，皆可純 HTTP，無需瀏覽器自動化）
==============================================================================

入口清單：https://www.land.moi.gov.tw/chhtml/landnoqry/51
地段代碼備援：https://lisp.land.moi.gov.tw/MMS/MMSpage.aspx#gobox02

【A. 標準 JSP 系統】臺中市/臺南市/屏東縣（多數縣市同款，URL 見 STANDARD_SYSTEMS）
  頁面：query/rwd/oldnew.jsp（RWD 版）
  流程：
    1. GET oldnew.jsp → 取 session cookie、頁內各端點專屬 csrf.param、
       SiteArea 下拉（行政區 → 代碼，如 潭子區=BH-17、仁德區=DI-32、高樹鄉=TB-11；
       前兩碼為地政事務所碼，同段藉.shp 段小段碼字首）。
    2. GET ../ajaxAreaR48.jsp?csrf.param=<頁內ajax專用csrf>&SiteArea=<代碼>
       &R48=R48&ViewActionName=oldnew_va&dt=<ms>
       ※ 注意路徑要上一層（rwd 版 JS 是 "./../ajaxAreaR48.jsp"），打錯層回 404。
       回傳格式：段名#,#段代碼#;#段名#,#段代碼#;#（尾 3 字元捨去）。
       此清單含「重測前舊段」（如高樹鄉同時列 東振段=1365 與 東振新段=0213）。
    3. POST oldnew.jsp?csrf.param=<form action 的 csrf>&...（照 form action）
       欄位：action=Query1, r=G, SiteArea, R48=段代碼(4位), R48check=同值,
             TypeCF=C(地號), NUM1=宗地母(4位), NUM2=宗地子(4位), TypeON=O 或 N
       ※ TypeON 兩方向都查再合併去重（實測輸入端會被比對到重測前或重測後任一側）。
    4. 結果表：|重測前|重測後| 標頭，之後每列 8 欄：
       行政區,段代碼,段小段名,舊地號(0000-0000),行政區,段代碼,段小段名,新地號。
       查無時頁面含「查無新舊地建號資料」。

【B. 桃園市自建 ASP.NET WebForms】
  入口 https://www.land.tycg.gov.tw/chaspx/SQry4.aspx/14 內嵌 iframe，
  實際查詢頁：https://www.land.tycg.gov.tw/chaspx/SQry4_frame.aspx
  流程（需帶 __VIEWSTATE/__EVENTVALIDATION，兩段式 postback）：
    1. GET SQry4_frame.aspx → 取 hidden 欄位。
    2. POST（newsecbtn=更新段別(Alt+Q)，帶 searchtype/xcity）→ 段別下拉換成該區
       清單（含重測前舊段，如 蘆竹區 坑子口段後壁厝小段(0046)）、新 viewstate。
       ※ 跳過此步直接查會回「錯誤訊息：請選擇段別」。
    3. POST（btnSearch1=確定查詢）欄位：searchtype(1=新查舊,2=舊查新)、
       xcity(01~13 行政區代碼)、xarea(4位段代碼)、xtype=1(地號)、
       xno=地號「0000-0000」格式。
    4. 結果 GridView1 表：|舊地段|舊地號|新地段|新地號|，查無時顯示「查無資料」。
  ※ 兩方向(searchtype 1/2)都查再合併。

【已驗證案例（2026-07-07）】
  桃園 蘆竹 坑子口段後壁厝小段 1-1   → 重測後 後壁段 12（舊查新）
  臺中 潭子 寶興段 176               → 重測前 聚興段新興小段 461
  屏東 高樹 東振段 1215              → 重測前 東振新段 707
  桃園 觀音 草漯段 2535-1、臺南 仁德 崑崙段 786-4 → 重測系統兩方向皆查無
  （母號 2535、786 皆存在於 111年2月地籍 → 疑為其後之逕為分割，需人工向地所確認）

【侷限】此類系統只涵蓋「地籍圖重測」對照；逕為分割/合併等異動查不到，
  查無者輸出「需人工處理」。其他縣市（同款 JSP 為多）遇到再於 STANDARD_SYSTEMS
  補 URL 即可；臺北市/高雄市等自建系統屆時另寫 adapter。
==============================================================================
"""
from __future__ import annotations

import re
import sys
import time
import warnings
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
import requests
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

warnings.filterwarnings("ignore")

REPO_ROOT = Path(__file__).resolve().parent.parent
UA = {"User-Agent": "Mozilla/5.0"}
TIMEOUT = 30

# 標準 JSP 系統（同款介面）的縣市 → oldnew.jsp URL
STANDARD_SYSTEMS = {
    "臺中市": "https://landquery.taichung.gov.tw/query/rwd/oldnew.jsp?menu=true&type=M&reqType=rwd",
    "臺南市": "https://land-query.tainan.gov.tw/query/rwd/oldnew.jsp?menu=false&type=R",
    "屏東縣": "https://www2.land.pthg.gov.tw/query/rwd/oldnew.jsp?menu=true&type=R&reqType=rwd",
}
TYCG_FRAME = "https://www.land.tycg.gov.tw/chaspx/SQry4_frame.aspx"


def full_section_name(seg: str, subseg: str | None) -> str:
    """組合完整段小段名：坑子口+後壁厝 → 坑子口段後壁厝小段。"""
    name = seg if seg.endswith("段") else seg + "段"
    if subseg and str(subseg) not in ("", "None", "nan"):
        sub = str(subseg)
        name += sub if sub.endswith("小段") else sub + "小段"
    return name


def fmt_no(mu: str, zi: str) -> str:
    """0331/0001 → 331-1；0176/0000 → 176。"""
    m, z = int(mu), int(zi)
    return f"{m}-{z}" if z else str(m)


# ---------------------------------------------------------------------------
# A. 標準 JSP 系統（臺中/臺南/屏東…）
# ---------------------------------------------------------------------------
class StandardJspSystem:
    def __init__(self, base_url: str):
        self.base = base_url
        self.s = requests.Session()
        self.s.headers.update(UA)
        r = self.s.get(base_url, timeout=TIMEOUT, verify=False)
        r.encoding = "utf-8"
        self.page = r.text
        m = re.search(r'<form action="([^"]+)"', self.page, re.I)
        self.post_url = urljoin(base_url, m.group(1))
        m = re.search(r"ajaxAreaR48\.jsp\?csrf\.param=([0-9A-F]+)", self.page)
        self.ajax_csrf = m.group(1) if m else ""

    def site_code(self, town: str) -> str | None:
        """行政區名 → SiteArea 代碼（如 潭子區 → BH-17）。"""
        m = re.search(r'<SELECT[^>]*NAME="SiteArea".*?</SELECT>', self.page, re.S | re.I)
        if not m:
            return None
        for v, name in re.findall(r'VALUE="([^"]+)"[^>]*>([^<]+)', m.group(0), re.I):
            if name.strip() == town or town in name:
                return v
        return None

    def section_list(self, site: str) -> list[tuple[str, str]]:
        """(段小段名, 段代碼) 清單。"""
        url = urljoin(
            self.base,
            f"../ajaxAreaR48.jsp?csrf.param={self.ajax_csrf}&SiteArea={site}"
            f"&R48=R48&ViewActionName=oldnew_va&dt={int(time.time() * 1000)}",
        )
        r = self.s.get(url, timeout=TIMEOUT, verify=False, headers={"Referer": self.base})
        r.encoding = "utf-8"
        res = r.text.strip()
        if "#,#" not in res:
            raise RuntimeError(f"段清單 ajax 失敗：{res[:120]}")
        out = []
        for it in res[:-3].split("#;#"):
            name, code = it.split("#,#")
            out.append((name.strip(), code.strip()))
        return out

    def query(self, site: str, r48: str, mu: str, zi: str) -> list[dict]:
        """雙向查詢合併，回傳 [{old_town,old_code,old_sec,old_no,new_town,new_code,new_sec,new_no}]"""
        rows, seen = [], set()
        for typeon in ("O", "N"):
            data = {
                "action": "Query1", "r": "G",
                "SiteArea": site, "R48": r48, "R48check": r48,
                "TypeCF": "C", "NUM1": mu, "NUM2": zi, "TypeON": typeon,
            }
            r = self.s.post(self.post_url, data=data, timeout=TIMEOUT, verify=False,
                            headers={"Referer": self.base})
            r.encoding = "utf-8"
            body = re.sub(r"<script.*?</script>", "", r.text, flags=re.S | re.I)
            for tb in re.findall(r"<table[^>]*>.*?</table>", body, re.S | re.I):
                if "重測前" not in tb:
                    continue
                cells = [re.sub(r"<[^>]+>|\s", "", c) for c in re.findall(r"<td[^>]*>(.*?)</td>", tb, re.S | re.I)]
                cells = [c for c in cells if c]
                # 標頭後每 8 欄一列：行政區,段代碼,段小段名,舊地號,行政區,段代碼,段小段名,新地號
                data_cells = [c for c in cells if c not in
                              ("重測前", "重測後", "行政區", "段代碼", "段名小段名", "段小段名", "舊地號", "新地號")]
                for i in range(0, len(data_cells) - 7, 8):
                    g = data_cells[i:i + 8]
                    key = tuple(g)
                    if key not in seen and re.match(r"\d{4}-\d{4}", g[3]):
                        seen.add(key)
                        rows.append(dict(zip(
                            ("old_town", "old_code", "old_sec", "old_no",
                             "new_town", "new_code", "new_sec", "new_no"), g)))
        return rows


# ---------------------------------------------------------------------------
# B. 桃園市 ASP.NET 系統
# ---------------------------------------------------------------------------
class TaoyuanSystem:
    def __init__(self):
        self.s = requests.Session()
        self.s.headers.update(UA)
        r = self.s.get(TYCG_FRAME, timeout=TIMEOUT, verify=False)
        r.encoding = "utf-8"
        self.page = r.text

    @staticmethod
    def _hiddens(t: str) -> dict:
        d = {}
        for name in ("__VIEWSTATE", "__VIEWSTATEGENERATOR", "__VIEWSTATEENCRYPTED", "__EVENTVALIDATION"):
            m = re.search(r'name="%s"[^>]*value="([^"]*)"' % name, t)
            d[name] = m.group(1) if m else ""
        return d

    def city_code(self, town: str) -> str | None:
        m = re.search(r'<select[^>]*name="xcity"[^>]*>(.*?)</select>', self.page, re.S | re.I)
        if not m:
            return None
        for v, name in re.findall(r'<option[^>]*value="([^"]*)"[^>]*>([^<]*)', m.group(1)):
            if town in name:
                return v
        return None

    def query(self, town: str, section_name: str, mu: str, zi: str) -> tuple[list[dict], str]:
        """雙向查詢合併。回傳 (rows, note)；rows 同標準系統欄位（town 用輸入值）。"""
        xcity = self.city_code(town)
        if not xcity:
            raise RuntimeError(f"桃園系統無此行政區：{town}")
        rows, seen = [], set()
        xarea_found = None
        for searchtype in ("2", "1"):  # 2=舊地號查新地號, 1=新地號查舊地號
            # step1 更新段別（取得該區段清單與新 viewstate）
            d1 = self._hiddens(self.page)
            d1.update({"searchtype": searchtype, "xcity": xcity, "xtype": "1",
                       "xno": "請輸入地建號", "newsecbtn": "更新段別(Alt+Q)"})
            r1 = self.s.post(TYCG_FRAME, data=d1, timeout=TIMEOUT, verify=False,
                             headers={"Referer": TYCG_FRAME})
            r1.encoding = "utf-8"
            m = re.search(r'<select[^>]*name="xarea"[^>]*>(.*?)</select>', r1.text, re.S | re.I)
            xarea = None
            for v, name in re.findall(r'<option[^>]*value="([^"]*)"[^>]*>([^<]*)', m.group(1) if m else ""):
                if re.sub(r"\(\d+\)$", "", name.strip()) == section_name:
                    xarea = v
                    break
            if not xarea:
                continue
            xarea_found = xarea
            # step2 查詢
            d2 = self._hiddens(r1.text)
            d2.update({"searchtype": searchtype, "xcity": xcity, "xarea": xarea,
                       "xtype": "1", "xno": f"{int(mu):04d}-{int(zi):04d}",
                       "btnSearch1": "確定查詢"})
            r2 = self.s.post(TYCG_FRAME, data=d2, timeout=TIMEOUT, verify=False,
                             headers={"Referer": TYCG_FRAME})
            r2.encoding = "utf-8"
            body = re.sub(r"<script.*?</script>", "", r2.text, flags=re.S | re.I)
            for tb in re.findall(r"<table[^>]*>.*?</table>", body, re.S | re.I):
                if "舊地段" not in tb:
                    continue
                cells = [re.sub(r"<[^>]+>|\s", "", c) for c in re.findall(r"<td[^>]*>(.*?)</td>", tb, re.S | re.I)]
                cells = [c for c in cells if c and c not in ("舊地段", "舊地號", "新地段", "新地號")]
                for i in range(0, len(cells) - 3, 4):
                    g = cells[i:i + 4]
                    key = tuple(g)
                    if key not in seen and re.match(r"\d{4}-\d{4}", g[1]):
                        seen.add(key)
                        rows.append({
                            "old_town": town, "old_code": "", "old_sec": g[0], "old_no": g[1],
                            "new_town": town, "new_code": "", "new_sec": g[2], "new_no": g[3],
                        })
        note = "" if xarea_found else f"桃園系統段別清單查無「{section_name}」"
        return rows, note


# ---------------------------------------------------------------------------
# 地籍 shp 交叉驗證（111年2月版）：查得對照是否存在／查無者母號是否存在
# ---------------------------------------------------------------------------
class ShpChecker:
    """惰性載入各縣市地籍屬性表（只讀段碼/宗地母/宗地子，不讀幾何）。"""

    def __init__(self):
        with (REPO_ROOT / "config" / "county_registry.yaml").open(encoding="utf-8") as f:
            reg = yaml.safe_load(f)
        self.reg = reg.get("counties", reg)
        self.sec_idx = pd.read_csv(REPO_ROOT / "config" / "section_index.csv",
                                   dtype=str, encoding="utf-8-sig")
        self._cache: dict[str, pd.DataFrame] = {}

    def _attrs(self, county: str) -> pd.DataFrame | None:
        key = county.replace("臺", "台")
        for k in (county, key, county.replace("台", "臺")):
            if k in self.reg and self.reg[k].get("cadastral"):
                path = REPO_ROOT / self.reg[k]["cadastral"]
                break
        else:
            return None
        if county not in self._cache:
            import pyogrio
            self._cache[county] = pyogrio.read_dataframe(
                str(path), read_geometry=False, columns=["段碼", "宗地母", "宗地子"])
        return self._cache[county]

    def section_code(self, county: str, town: str, sec_name: str) -> str | None:
        """完整段小段名 → 4位段碼（用 section_index.csv）。"""
        df = self.sec_idx
        cand = df[df["縣市"].str.replace("台", "臺") == county.replace("台", "臺")]
        for _, r in cand.iterrows():
            if full_section_name(str(r["段"]), None if pd.isna(r["小段"]) else r["小段"]) == sec_name:
                return r["段碼"]
        return None

    def exists(self, county: str, town: str, sec_name: str, mu: int, zi: int) -> bool | None:
        """該宗地是否存在於地籍 shp。None=無法判斷（缺段碼或縣市檔）。"""
        code = self.section_code(county, town, sec_name)
        if not code:
            return None
        attrs = self._attrs(county)
        if attrs is None:
            return None
        hit = attrs[(attrs["段碼"] == code) & (attrs["宗地母"] == f"{mu:04d}")
                    & (attrs["宗地子"] == f"{zi:04d}")]
        return len(hit) > 0


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
HEADERS = ["案場ID", "縣市", "鄉鎮區", "原段/小段", "原地號", "查得新段名",
           "查得新地號", "查詢狀態", "查詢來源", "備註"]


def normalize_county(name: str) -> str:
    return str(name).replace("台", "臺")


def process_row(row: dict, jsp_cache: dict, tycg: list, checker: ShpChecker | None) -> dict:
    county = normalize_county(row["縣市"])
    town = str(row["鄉鎮區"])
    sec_name = full_section_name(str(row["段"]), row.get("小段"))
    mu, zi = str(row["宗地母"]), str(row["宗地子"])
    out = {
        "案場ID": row["案場ID"], "縣市": row["縣市"], "鄉鎮區": town,
        "原段/小段": sec_name, "原地號": fmt_no(mu, zi),
        "查得新段名": "", "查得新地號": "", "查詢狀態": "", "查詢來源": "", "備註": "",
    }
    try:
        if county == "桃園市":
            out["查詢來源"] = "https://www.land.tycg.gov.tw/chaspx/SQry4.aspx/14"
            if not tycg:
                tycg.append(TaoyuanSystem())
            rows, note = tycg[0].query(town, sec_name, mu, zi)
            matches = rows
        elif county in STANDARD_SYSTEMS:
            out["查詢來源"] = STANDARD_SYSTEMS[county]
            if county not in jsp_cache:
                jsp_cache[county] = StandardJspSystem(STANDARD_SYSTEMS[county])
            sys_ = jsp_cache[county]
            site = sys_.site_code(town)
            if not site:
                raise RuntimeError(f"查無行政區代碼：{town}")
            sections = sys_.section_list(site)
            code = next((c for n, c in sections if n == sec_name), None)
            if not code:
                out["查詢狀態"] = "查無"
                out["備註"] = f"系統段別清單查無「{sec_name}」，需人工處理"
                return out
            matches = sys_.query(site, code, mu, zi)
            note = ""
        else:
            out["查詢狀態"] = "系統異常"
            out["備註"] = f"尚未支援 {county} 之新舊地號系統（遇到再擴充），需人工處理"
            return out

        if not matches:
            out["查詢狀態"] = "查無"
            remark = ("重測新舊地號系統雙向查無此地號；"
                      "疑為重測以外之異動（如逕為分割/合併），需人工處理" + ("；" + note if note else ""))
            # 交叉驗證：母號是否存在於 111年2月地籍（佐證「其後分割」假說）
            if checker and int(zi) != 0:
                parent = checker.exists(county, town, sec_name, int(mu), 0)
                if parent:
                    remark += f"（母號 {int(mu)} 存在於111年2月地籍，疑為其後自母號分割）"
            out["備註"] = remark
            return out

        # 判斷輸入是被比對到「舊」側或「新」側，輸出對側（即查得之對照段/地號）
        got_secs, got_nos, remarks = [], [], []
        for r in matches:
            in_no = f"{int(mu):04d}-{int(zi):04d}"
            if r["old_sec"] == sec_name and r["old_no"] == in_no:
                got, direction = (r["new_sec"], r["new_no"]), "原地號為重測前舊地號，查得為重測後新段名/新地號"
            else:
                got, direction = (r["old_sec"], r["old_no"]), ("原地號為重測後新地號，查得為重測前舊段名/舊地號"
                                                               "（111年2月版地籍以舊段名收錄，可據此回填比對）")
            g_sec, g_no = got
            g_mu, g_zi = (int(x) for x in g_no.split("-"))
            got_secs.append(g_sec)
            got_nos.append(fmt_no(*g_no.split("-")))
            remarks.append(direction)
            if checker:
                ok = checker.exists(county, town, g_sec, g_mu, g_zi)
                if ok is True:
                    remarks.append("已驗證：查得對照地號存在於111年2月地籍")
                elif ok is False:
                    remarks.append("注意：查得對照地號不存在於111年2月地籍，請人工確認")
        out["查得新段名"] = "、".join(dict.fromkeys(got_secs))
        out["查得新地號"] = "、".join(dict.fromkeys(got_nos))
        out["查詢狀態"] = "成功"
        out["備註"] = "；".join(dict.fromkeys(remarks))
    except Exception as exc:  # noqa: BLE001
        out["查詢狀態"] = "系統異常"
        out["備註"] = f"{exc}；需人工處理"
    return out


def main() -> None:
    case = sys.argv[1] if len(sys.argv) > 1 else "20260707-清潔隊部資訊"
    outdir = REPO_ROOT / "outputs" / case
    src = outdir / "03_無法對應清單.xlsx"
    dst = outdir / "04_新舊地號對照清單.xlsx"

    ws = load_workbook(src, read_only=True).active
    it = ws.iter_rows(values_only=True)
    cols = [str(c) for c in next(it)]
    rows = [dict(zip(cols, r)) for r in it if any(v is not None for v in r)]
    print(f"讀入 {src.name}：{len(rows)} 筆")

    try:
        checker: ShpChecker | None = ShpChecker()
    except Exception as exc:  # noqa: BLE001
        print(f"[警告] 地籍交叉驗證停用（{exc}）")
        checker = None
    jsp_cache: dict = {}
    tycg: list = []
    results = []
    for row in rows:
        res = process_row(row, jsp_cache, tycg, checker)
        print(f"  [{res['查詢狀態']}] {res['案場ID']} {res['縣市']}{res['鄉鎮區']} "
              f"{res['原段/小段']} {res['原地號']} → {res['查得新段名']} {res['查得新地號']}")
        results.append(res)
        time.sleep(1)  # 禮貌間隔

    wb = Workbook()
    w = wb.active
    w.title = "新舊地號對照清單"
    w.append(HEADERS)
    for c in w[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for res in results:
        w.append([res[h] for h in HEADERS])
        if res["查詢狀態"] != "成功":
            for c in w[w.max_row]:
                c.fill = yellow
    widths = [8, 8, 10, 22, 10, 22, 14, 10, 46, 60]
    for i, wd in enumerate(widths, 1):
        w.column_dimensions[w.cell(row=1, column=i).column_letter].width = wd
    for r in w.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(dst)
    n_ok = sum(1 for r in results if r["查詢狀態"] == "成功")
    print(f"輸出 {dst}（成功 {n_ok}/{len(results)}）")


if __name__ == "__main__":
    main()
