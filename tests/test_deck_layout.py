"""簡報產出防線：缺圖、markdown、手改鎖與正常存檔。只寫 tmp_path。"""
from pathlib import Path
import sys

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import deck_layout as deck


def test_missing_picture_blocks_save(tmp_path):
    deck.init(str(tmp_path), 'test.pptx')
    deck.pic(deck.new(), 'missing.png', 0, 0, Inches(1))
    with pytest.raises(SystemExit, match='簡報缺'):
        deck.save()
    assert not (tmp_path / 'test.pptx').exists()


@pytest.mark.parametrize('text', ['**粗體**', '## 標題', '`程式`'])
@pytest.mark.parametrize('in_table', [False, True])
def test_markdown_blocks_save(tmp_path, text, in_table):
    deck.init(str(tmp_path), 'test.pptx')
    slide = deck.new()
    if in_table:
        deck.table(slide, 0, 0, Inches(5), [[text]])
    else:
        deck.tb(slide, 0, 0, Inches(5), Inches(1), text)
    with pytest.raises(SystemExit, match='markdown'):
        deck.save()
    assert not (tmp_path / 'test.pptx').exists()


def test_picture_success_resets_missing_state_and_protects_manual_edits(tmp_path):
    deck.init(str(tmp_path), 'unused.pptx')
    deck.pic(deck.new(), 'missing.png', 0, 0, Inches(1))
    deck.init(str(tmp_path), 'test.pptx')
    Image.new('RGB', (40, 20), 'green').save(tmp_path / 'map.png')
    slide = deck.new()
    deck.pic(slide, 'map.png', 0, 0, Inches(1))
    deck.tb(slide, 0, Inches(2), Inches(5), Inches(1), '正常文字')
    deck.save()
    assert (tmp_path / 'test.pptx.dump.txt').is_file()
    prs = Presentation(tmp_path / 'test.pptx')
    assert len(prs.slides) == 1
    prs.slides[0].shapes[-1].text = '人工修改'
    prs.save(tmp_path / 'test.pptx')
    with pytest.raises(SystemExit, match='手改'):
        deck.init(str(tmp_path), 'test.pptx')
